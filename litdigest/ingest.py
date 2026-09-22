"""Stage 1: read the spreadsheet into per-paper records, preserving your Notes column."""
import pandas as pd
from openpyxl import load_workbook

from . import config, store


def header_row(ws, limit: int = 12) -> int | None:
    """1-based index of the row holding the column names, found by looking for Title."""
    for r in range(1, min(limit, ws.max_row) + 1):
        for c in range(1, ws.max_column + 1):
            if str(ws.cell(r, c).value).strip().lower() == "title":
                return r
    return None


def run() -> int:
    wb = load_workbook(config.SOURCE_XLSX)          # styles are needed for the star
    ws = wb[wb.sheetnames[0]]
    head = header_row(ws)
    if head is None:
        raise SystemExit(f"no 'Title' column found in {config.SOURCE_XLSX}")
    df = pd.read_excel(config.SOURCE_XLSX, header=head - 1)
    cols = [c for c in df.columns if not str(c).startswith("Unnamed")]
    df = df[cols]
    df.columns = [str(c).strip().lower() for c in df.columns]
    df = df.dropna(subset=["title"])

    from .sheet import columns, is_starred
    cols = columns(ws)
    stars = {}
    if cols.get("num") and cols.get("title"):
        for r in range(head + 1, ws.max_row + 1):
            try:
                stars[int(ws.cell(r, cols["num"]).value)] = is_starred(
                    ws.cell(r, cols["title"]))
            except (TypeError, ValueError):
                continue

    n = 0
    for i, row in enumerate(df.itertuples(index=False), start=1):
        num = int(row.num) if pd.notna(getattr(row, "num", None)) else i
        title = str(row.title).strip()
        notes = getattr(row, "notes", None)
        notes = str(notes).strip() if pd.notna(notes) else ""

        def sync(rec):
            if rec.get("title") and rec["title"] != title:
                # The row now describes a different paper, so everything derived from
                # the old one is wrong. The note and the star are not listed here: they
                # live in the spreadsheet and are re-read from it on every ingest.
                for k in ("arxiv", "text", "glance", "deep", "topic",
                          "equations", "macros", "figures", "chat", "errors"):
                    rec.pop(k, None)
            rec.update(title=title, notes=notes, starred=stars.get(num, False))
        store.update(num, sync)
        n += 1
    return n
