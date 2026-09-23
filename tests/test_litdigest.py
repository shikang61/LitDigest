"""Tests for the parts that are easy to break and awkward to notice: the title
matcher, the tagged-section parser, header detection, the spreadsheet reader, the
record store, and the launcher's health probe.

No network and no model calls -- everything here runs offline.
"""
import sys
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from litdigest import arxiv, generate, ingest  # noqa: E402


# --- title matching -------------------------------------------------------

def test_normalise_strips_punctuation_and_case():
    assert arxiv.normalise("Grad--Shafranov: Shape Effects!") == "grad shafranov shape effects"


def test_exact_title_scores_one():
    t = "Entropy stable finite difference schemes"
    ratio, overlap = arxiv._score(arxiv.normalise(t), t)
    assert ratio == pytest.approx(1.0)
    assert overlap == pytest.approx(1.0)


def test_retitled_paper_is_caught_by_word_overlap():
    """arXiv v3 reordered the title; character similarity alone would miss it."""
    sheet = arxiv.normalise(
        "Analytical Grad-Shafranov equilibria for prescribed boundaries: shape effects")
    live = "Shape effects and Shafranov shift reversal in analytical Grad--Shafranov equilibria"
    ratio, overlap = arxiv._score(sheet, live)
    assert ratio < 0.72, "character similarity is not what rescues this"
    assert overlap >= 0.6


def test_unrelated_paper_does_not_match():
    sheet = arxiv.normalise("Entropy stable schemes for the Euler equations")
    ratio, overlap = arxiv._score(sheet, "Deep learning for protein folding")
    assert ratio < 0.72 and overlap < 0.75


def test_stopwords_excluded_from_overlap():
    assert "the" not in arxiv._content("the theory of the thing")
    assert "theory" in arxiv._content("the theory of the thing")


# --- the tagged-section format --------------------------------------------

GLANCE = """[CLAIM]
A claim in one line.
[SCORE]
2
[VERDICT]
Not worth it. Read section 3.
[TRICK]
The mechanism.
[HOLDS UP]
One caveat.
"""


def test_parse_splits_every_section():
    out = generate.parse(GLANCE)
    assert out["claim"] == "A claim in one line."
    assert out["score"] == 2
    assert out["holds_up"] == "One caveat."


def test_parse_handles_a_multi_word_tag():
    assert "holds_up" in generate.parse(GLANCE)


def test_parse_of_a_half_written_stream():
    """The panel renders while the model is still writing, so partial text must parse."""
    partial = GLANCE[:GLANCE.index("[TRICK]")]
    out = generate.parse(partial)
    assert out["claim"] and out["score"] == 2
    assert "trick" not in out


def test_parse_survives_a_missing_score():
    out = generate.parse("[CLAIM]\nx\n[SCORE]\n\n[VERDICT]\ny\n")
    assert out["score"] is None


def test_body_text_is_not_mistaken_for_a_tag():
    out = generate.parse("[CLAIM]\nSee [EQ] in the paper.\n[SCORE]\n3\n")
    assert "See [EQ] in the paper." == out["claim"]


# --- the spreadsheet reader -----------------------------------------------

def _sheet(tmp_path, header_row=3):
    wb = Workbook()
    ws = wb.active
    ws.title = "LitFeed"
    for i, name in enumerate(("Num", "Title", "Notes"), start=2):
        ws.cell(header_row, i, name)
    ws.cell(header_row + 1, 2, 1)
    ws.cell(header_row + 1, 3, "A Paper About Things")
    ws.cell(header_row + 1, 4, "my note")
    path = tmp_path / "sheet.xlsx"
    wb.save(path)
    return path


def test_header_row_is_found_below_blank_rows(tmp_path):
    ws = load_workbook(_sheet(tmp_path))["LitFeed"]
    assert ingest.header_row(ws) == 3


def test_header_row_found_at_the_top(tmp_path):
    ws = load_workbook(_sheet(tmp_path, header_row=1))["LitFeed"]
    assert ingest.header_row(ws) == 1


def test_header_row_missing_returns_none(tmp_path):
    wb = Workbook(); wb.active["B2"] = "nothing useful"
    p = tmp_path / "bad.xlsx"; wb.save(p)
    assert ingest.header_row(load_workbook(p).active) is None


def test_ingest_reads_titles_and_notes(tmp_path, monkeypatch):
    from litdigest import config, store
    monkeypatch.setattr(config, "SOURCE_XLSX", _sheet(tmp_path))
    monkeypatch.setattr(config, "PAPER_DIR", tmp_path / "papers")
    (tmp_path / "papers").mkdir()
    assert ingest.run() == 1
    rec = store.load(1)
    assert rec["title"] == "A Paper About Things"
    assert rec["notes"] == "my note"


def test_changing_a_title_discards_its_cached_work(tmp_path, monkeypatch):
    from litdigest import config, store
    monkeypatch.setattr(config, "SOURCE_XLSX", _sheet(tmp_path))
    monkeypatch.setattr(config, "PAPER_DIR", tmp_path / "papers")
    (tmp_path / "papers").mkdir()
    ingest.run()
    rec = store.load(1)
    rec.update(arxiv={"id": "x"}, glance={"claim": "stale"})
    store.save(rec)

    wb = load_workbook(config.SOURCE_XLSX)
    wb["LitFeed"].cell(4, 3, "A Different Paper")
    wb.save(config.SOURCE_XLSX)
    ingest.run()

    after = store.load(1)
    assert after["title"] == "A Different Paper"
    assert "arxiv" not in after and "glance" not in after


# --- the record store -----------------------------------------------------

def test_save_leaves_no_temporary_file_behind(tmp_path, monkeypatch):
    from litdigest import config, store
    monkeypatch.setattr(config, "PAPER_DIR", tmp_path)
    store.save({"num": 1, "title": "x"})
    assert store.load(1)["title"] == "x"
    assert [p.name for p in tmp_path.iterdir()] == ["0001.json"]


def test_a_failed_save_does_not_destroy_the_previous_record(tmp_path, monkeypatch):
    """The old write truncated in place, so a crash mid-write cost the record --
    and one unreadable file makes every later read of the library fail."""
    from litdigest import config, store
    monkeypatch.setattr(config, "PAPER_DIR", tmp_path)
    store.save({"num": 1, "glance": {"claim": "the good one"}})

    # the write dies halfway through, which is what an interrupted one looks like
    real = Path.write_text

    def half(self, text, *a, **k):
        real(self, text[:len(text) // 2])
        raise OSError("disk full")
    Path.write_text = half
    try:
        with pytest.raises(OSError):
            store.save({"num": 1, "glance": {"claim": "the half-written one"}})
    finally:
        Path.write_text = real

    assert store.load(1)["glance"]["claim"] == "the good one"
    assert list(store.all_records())                      # still readable
    assert [p.name for p in tmp_path.iterdir()] == ["0001.json"]


def test_concurrent_saves_are_never_seen_half_written(tmp_path, monkeypatch):
    """Two stages can write the same paper at once -- a generation finishing while
    /api/papers re-reads the spreadsheet. A reader must see one of them, not both."""
    import json
    import threading
    from litdigest import config, store
    monkeypatch.setattr(config, "PAPER_DIR", tmp_path)
    store.save({"num": 1, "pad": "x"})

    stop = threading.Event()
    torn = []

    def write(tag):
        for _ in range(40):
            store.save({"num": 1, "tag": tag, "pad": tag * 20000})

    def read():
        while not stop.is_set():
            try:
                rec = store.load(1)
            except json.JSONDecodeError as exc:
                torn.append(str(exc))
                return
            if rec and len(rec.get("pad", "")) not in (1, 20000):
                torn.append("mixed payload")
                return

    readers = [threading.Thread(target=read) for _ in range(2)]
    writers = [threading.Thread(target=write, args=(c,)) for c in "ab"]
    for t in readers + writers:
        t.start()
    for t in writers:
        t.join()
    stop.set()
    for t in readers:
        t.join()

    assert not torn, torn
    assert not list(tmp_path.glob(".*tmp"))


# --- the launcher's probe -------------------------------------------------

def test_health_answers_without_reading_the_spreadsheet(monkeypatch):
    """launch.sh polls this twice a second while the server comes up. /api/papers
    re-reads the sheet and can run model calls, so the probe must not be that."""
    from starlette.testclient import TestClient
    import server

    def explode(*a, **k):
        raise AssertionError("the probe must not touch the spreadsheet")
    monkeypatch.setattr(server.ingest, "run", explode)

    with TestClient(server.app) as client:
        r = client.get("/api/health")
        assert r.status_code == 200
        assert r.json() == {"ok": True, "stale": False, "running": 0}
        with pytest.raises(AssertionError):
            client.get("/api/papers")          # the expensive one, for contrast


def test_health_says_when_the_code_changed_after_start(monkeypatch):
    """A server started before a fix was saved keeps running the old code, and the
    launcher used to hand over to it. This is how it knows to restart instead."""
    from starlette.testclient import TestClient
    import server
    monkeypatch.setattr(server, "STARTED_CODE", 0.0)
    with TestClient(server.app) as client:
        assert client.get("/api/health").json()["stale"] is True


def test_a_note_saved_during_deep_prep_survives(tmp_path, monkeypatch):
    """Clicking Go deeper blurs the note box, so the note and the deep read reach the
    server together. Deep prep used to load the record, spend seconds downloading
    the LaTeX source and figures, then save what it loaded -- undoing the note."""
    import server
    from litdigest import config, figures, sheet, store
    monkeypatch.setattr(config, "PAPER_DIR", tmp_path)
    monkeypatch.setattr(sheet, "set_note", lambda num, text: True)
    store.save({"num": 1, "title": "t", "notes": "", "text": {"intro": "x", "chars": 1},
                "arxiv": {"id": "2501.00001", "match_status": "ok"}, "equations": []})

    def slow_download(aid):
        server.note(1, text="typed meanwhile")       # lands mid-download
        return [{"file": "f.png"}]
    monkeypatch.setattr(figures, "extract", slow_download)

    server._prep_deep(1)
    rec = store.load(1)
    assert rec["notes"] == "typed meanwhile"
    assert rec["figures"] == [{"file": "f.png"}]


def test_deleting_a_note_clears_the_spreadsheet_cell(tmp_path, monkeypatch):
    from litdigest import config, sheet
    wb = Workbook()
    ws = wb.active
    ws.append(["Num", "Title", "Notes"])
    ws.append([1, "A paper", "old note"])
    path = tmp_path / "s.xlsx"
    wb.save(path)
    monkeypatch.setattr(config, "SOURCE_XLSX", path)
    monkeypatch.setattr(config, "BACKUP_DIR", tmp_path / "bk")

    assert sheet.set_note(1, "")
    assert load_workbook(path).active.cell(2, 3).value is None


def test_a_note_cannot_be_written_without_a_notes_column(tmp_path, monkeypatch):
    from litdigest import config, sheet
    wb = Workbook()
    ws = wb.active
    ws.append(["Num", "Title"])
    ws.append([1, "A paper"])
    path = tmp_path / "s.xlsx"
    wb.save(path)
    monkeypatch.setattr(config, "SOURCE_XLSX", path)
    monkeypatch.setattr(config, "BACKUP_DIR", tmp_path / "bk")

    assert sheet.set_note(1, "a note") is False


# --- the prompts ------------------------------------------------------------

def test_the_deep_prompt_differs_from_the_glance_prompt():
    """DEEP_SYS is GLANCE_SYS with passages swapped by str.replace, which does
    nothing at all once a passage is reworded."""
    assert generate.DEEP_SYS != generate.GLANCE_SYS
    assert "Bold for skimming" in generate.GLANCE_SYS
    assert "Bold for skimming" not in generate.DEEP_SYS


def test_the_time_estimate_follows_the_latest_runs(monkeypatch):
    """Five slow runs after five fast ones: the estimate is the slow time."""
    from litdigest import store
    runs = [{"deep": {"seconds": s, "generated_at": f"2026-09-{d:02d}T10:00:00"}}
            for d, s in enumerate([180] * 5 + [430] * 5, start=1)]
    monkeypatch.setattr(store, "all_records", lambda: iter(runs[::-1]))
    assert generate.expected_seconds("deep") == 430


def test_an_overlong_paper_keeps_its_start_and_end(monkeypatch):
    from litdigest import config, extract
    monkeypatch.setattr(config, "FULLTEXT_CHARS", 1000)
    monkeypatch.setattr(config, "EXCERPT_CHARS", 100)
    monkeypatch.setattr(extract, "pdf_text", lambda rec: "S" * 5000 + "E" * 100)

    text = generate._whole_paper({"title": "t", "arxiv": {}})
    body = text.split("FULL TEXT\n", 1)[1]
    assert body.startswith("S" * 900) and body.endswith("E" * 100)
    assert len(body) < 1100


# --- the web search -----------------------------------------------------------

def _drain(gen):
    """Everything a generator yields, and what it returns."""
    out = []
    try:
        while True:
            out.append(next(gen))
    except StopIteration as done:
        return out, done.value


def _fake_search(monkeypatch, events):
    from types import SimpleNamespace as NS

    from litdigest import llm
    monkeypatch.setattr(llm, "client",
                        lambda: NS(responses=NS(create=lambda **kw: iter(events))))


def test_web_search_shows_each_query_and_returns_the_notes(monkeypatch):
    from types import SimpleNamespace as NS

    from litdigest import llm
    _fake_search(monkeypatch, [
        NS(type="response.output_item.done",
           item=NS(type="web_search_call", action=NS(query="who cites it"))),
        NS(type="response.output_text.delta", delta="Two later "),
        NS(type="response.output_text.delta", delta="papers cite it."),
        NS(type="response.completed"),
    ])
    said, notes = _drain(llm.web_search("sys", "user"))
    assert said == [("think", ' Searching the web for "who cites it". ')]
    assert notes == "Two later papers cite it."


def test_notes_written_without_searching_are_dropped(monkeypatch):
    """They would reach the reading labelled as web notes, but are only the fast
    model's opinion."""
    from types import SimpleNamespace as NS

    from litdigest import llm
    _fake_search(monkeypatch, [NS(type="response.output_text.delta", delta="I recall...")])
    assert _drain(llm.web_search("sys", "user")) == ([], "")


def _fake_reading(monkeypatch, search):
    from litdigest import extract, llm
    seen = {}

    def reading(system, user, max_tokens):
        seen["user"] = user
        yield "say", "[CLAIM]\nx"
    monkeypatch.setattr(llm, "web_search", search)
    monkeypatch.setattr(llm, "stream_parts", reading)
    monkeypatch.setattr(extract, "pdf_text", lambda rec: "the paper")
    return seen


def test_what_the_search_found_reaches_the_reading(monkeypatch):
    def search(system, user):
        yield "think", " Searching the web for x. "
        return "A 2025 follow-up disputes it."
    seen = _fake_reading(monkeypatch, search)

    parts = list(generate.stream_glance({"title": "t", "arxiv": {}}))
    assert parts == [("think", " Searching the web for x. "), ("say", "[CLAIM]\nx")]
    assert "A 2025 follow-up disputes it." in seen["user"]


def test_a_failed_web_search_still_reads_the_paper(monkeypatch):
    def search(system, user):
        raise RuntimeError("search is down")
        yield
    seen = _fake_reading(monkeypatch, search)

    parts = list(generate.stream_glance({"title": "t", "arxiv": {}}))
    assert parts[0][0] == "think" and "failed" in parts[0][1]
    assert parts[-1] == ("say", "[CLAIM]\nx")
    assert "WEB SEARCH NOTES" not in seen["user"]


# --- generations outlive the page -------------------------------------------

def _fake_generation(tmp_path, monkeypatch, gate=None):
    import server
    from litdigest import config, store
    monkeypatch.setattr(config, "PAPER_DIR", tmp_path)
    monkeypatch.setattr(generate, "expected_seconds", lambda key: 1.0)
    store.save({"num": 1, "title": "t"})
    calls = []

    def parts(rec):
        calls.append(1)
        if gate is not None:
            gate.wait(5)
        yield "say", "[KEY POINTS]\n- a point\n"
    return server, store, parts, calls


def _wait_idle(server):
    import time
    for _ in range(200):
        if not server.RUNS:
            return
        time.sleep(0.01)
    raise AssertionError("the generation never finished")


def test_a_generation_is_saved_with_nobody_reading_it(tmp_path, monkeypatch):
    """A deep read takes minutes. It used to run inside the browser's request, so
    closing the tab or quitting the page before the end threw it away."""
    server, store, parts, _ = _fake_generation(tmp_path, monkeypatch)
    server._events(parts, 1, "glance")             # the response is never read
    _wait_idle(server)
    assert store.load(1)["glance"]["key_points"] == "- a point"


def test_asking_again_follows_the_run_instead_of_paying_twice(tmp_path, monkeypatch):
    import threading
    gate = threading.Event()
    server, store, parts, calls = _fake_generation(tmp_path, monkeypatch, gate)
    first = server._events(parts, 1, "glance")
    second = server._events(parts, 1, "glance")    # the card reopened, the page reloaded
    assert first is not second and len(server.RUNS) == 1
    gate.set()
    _wait_idle(server)
    assert calls == [1]
